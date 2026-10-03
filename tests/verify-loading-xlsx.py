import openpyxl, zipfile
p="artifacts/loading-list.xlsx"
assert zipfile.ZipFile(p).testzip() is None
s=openpyxl.load_workbook(p).active
rows=[r for r in s.iter_rows(min_row=5,values_only=True) if r[6]=="Otomatik H adedi"]
assert {r[8]:r[5] for r in rows}=={"tuna-23":10,"tuna-24":10,"tuna-25":9,"tuna-26":2,"tuna-27":1,"tuna-28":4}
assert sum(r[5] for r in rows)==36
assert s.freeze_panes=="A5"
assert all(c.data_type!="f" for row in s for c in row)
print("PASS XLSX structure, six Tuna matches, numeric quantities and text cells")

corners=[r for r in s.iter_rows(min_row=5,values_only=True) if r[1]=="Köşe direği"]
assert len(corners)==1 and corners[0][2]=="98 × 98 × 2500 mm" and corners[0][5]==8 and corners[0][8]=="tuna-22"
print("PASS XLSX corner product 98x98x2500, quantity 8 and Tuna reference")

us=[r for r in s.iter_rows(min_row=5,values_only=True) if r[1]=="Çektirme U"]
assert len(us)==1 and us[0][2]=="60 × 2440 mm" and us[0][4]==3 and us[0][5]==4 and us[0][10]==1 and us[0][8]=="tuna-29"
print("PASS XLSX U 60x2440: 3 installed + 1 spare = 4 shipped")

frames=[r for r in s.iter_rows(min_row=5,values_only=True) if r[1]=="Alt çerçeve"]
assert {r[2]:(r[4],r[5],r[10]) for r in frames}=={"100 × 2500 mm":(15,16,1),"60 × 2500 mm":(9,10,1)}
print("PASS XLSX bottom frames: base, shipment and spare")

assert s['F4'].value=='Toplam sevk (yedek dahil)'
assert s['K4'].value=='Yedek'
assert s['N4'].value=='Adet'
assert all(r[14]=='Yedek toplam sevke dahil; tekrar eklemeyin' for r in frames+us)
assert {r[2]:r[13] for r in frames}=={'100 × 2500 mm':15,'60 × 2500 mm':9}
manual=[r for r in s.iter_rows(min_row=5,values_only=True) if r[6]=='Manuel doğrulandı']
assert all(r[10] in (None, '') and 'doğrulanmadı' in r[14] for r in manual)
print('PASS XLSX explicit spare inclusion and manual uncertainty')

omegas=[r for r in s.iter_rows(min_row=5,values_only=True) if r[1]=='Duvar omegası' and r[2]=='60 × 2500 mm']
assert len(omegas)==1 and omegas[0][2]=='60 × 2500 mm'
assert (omegas[0][5],omegas[0][10],omegas[0][13])==(9,0,9)
print('PASS XLSX interior omega 9 stock pieces, no spare')

allrows=list(s.iter_rows(min_row=5,values_only=True))
assert {r[1]:r[5] for r in allrows if r[1] in ('Baş makas omegası','Baş makas Z sacı','Saçak omegası')}=={'Baş makas omegası':8,'Baş makas Z sacı':8,'Saçak omegası':7}
assert sum(1 for r in allrows if r[8]=='tuna-41')==1
assert [s.cell(4,i).value for i in (16,17,18)]==['Net boy (mm)','Kesim payı (mm)','Sevk boyu (mm)']
assert sorted((r[15],r[16],r[17]) for r in allrows if r[1]=='Veranda kirişi')==[(645,100,745),(2527.5,100,2627.5),(4970,100,5070)]
print('PASS XLSX separate head/Z, eaves, and net/allowance/shipping columns')

sheets=[r for r in allrows if r[1]=='Saçak sacı']
assert len(sheets)==1 and sheets[0][2]=='300 × 2800 mm' and sheets[0][5]==7 and sheets[0][8]=='tuna-35'
assert '300 mm bindirme' in sheets[0][7]
print('PASS XLSX separate eaves sheet with overlap explanation')
